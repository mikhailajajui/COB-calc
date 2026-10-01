#!/usr/bin/env python3
"""REF-01 extraction: the saved calculation in the real workbook (QA, T5, 2026-09-27).

Reads the CACHED values (openpyxl data_only=True, read_only=True) that Excel stored the
last time the workbook was calculated and saved. Nothing is recomputed and the workbook is
never written. Output: ca_ref01_workbook_saved.json next to this script.

Usage:
  python3 extract_ref01.py [path/to/Cost of Borrowing Rate Calc_Current.xlsm]
Default path: ../../../../Cost of Borrowing Rate Calc_Current.xlsm (the cob/ folder).
"""
import datetime
import hashlib
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT = os.path.normpath(os.path.join(HERE, '..', '..', '..', '..', 'Cost of Borrowing Rate Calc_Current.xlsm'))
SHEET = 'Calculator'
FIRST_ROW = 37
ROW_COLS = [  # column -> (fixture key, workbook header in row 36)
    ('B', 'n', 'Pymt #'),
    ('C', 'date', 'Payment Date'),
    ('D', 'open_loan', 'Loan (opening)'),
    ('E', 'open_principal', 'Principle (opening)'),
    ('F', 'open_fees', 'Fees (opening)'),
    ('G', 'new_int', 'New Int'),
    ('H', 'total_int', 'Total Int (cumulative accrued)'),
    ('I', 'payment', 'Payment'),
    ('J', 'interest_paid', 'Interest Paid'),
    ('K', 'fees_paid', 'Fees Paid'),
    ('L', 'principal_paid', 'Principle Paid'),
    ('M', 'close_fees', 'Fees (closing)'),
    ('N', 'close_principal', 'Principle (closing)'),
    ('O', 'close_loan', 'Loan (closing)'),
]
INPUT_CELLS = {
    'D4': 'loan_amount', 'D6': 'financed_fees', 'D8': 'non_financed_fees',
    'D10': 'calculated_rate_pct', 'F10': 'f10_rate_entry', 'G10': 'rate_type',
    'D12': 'payment_frequency', 'D14': 'disbursal_date', 'D16': 'term_years',
    'D18': 'term_months', 'D22': 'first_payment_date', 'D24': 'end_date', 'D26': 'payment_amount',
}
SUMMARY_CELLS = {
    'D29': 'cob_amount', 'D31': 'cob_rate_pct', 'G29': 'total_payment', 'G31': 'n_payments',
    'J29': 'total_interest', 'J31': 'principal_paid', 'M29': 'trigger_rate_pct',
    'D20': 'term_days', 'I193': 'payment_column_sum',
}
CONVERTER_CELLS = {'B4': 'mortgage_rate_pct', 'B5': 'compounding_per_year', 'C5': 'compounding_label'}


def norm(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    return v


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    with open(path, 'rb') as fh:
        sha = hashlib.sha256(fh.read()).hexdigest()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET]
    cv = wb['Semi-Annual Rate Converter']

    headers = {c: norm(ws[f'{c}36'].value) for c, _, _ in ROW_COLS}
    rows = []
    r = FIRST_ROW
    while ws[f'B{r}'].value is not None:
        rows.append({k: norm(ws[f'{c}{r}'].value) for c, k, _ in ROW_COLS})
        r += 1

    out = {
        'id': 'REF-01',
        'provenance': {
            'file': os.path.basename(path),
            'sha256': sha,
            'file_mtime_utc': datetime.datetime.fromtimestamp(os.path.getmtime(path), datetime.timezone.utc).isoformat(),
            'sheet': SHEET,
            'cells': {
                'inputs': sorted(INPUT_CELLS),
                'summary': sorted(SUMMARY_CELLS),
                'rows': f'B{FIRST_ROW}:O{r - 1}',
                'rate_source': "'Semi-Annual Rate Converter'!B4 (G10 = SEMI-ANNUAL, so D10 uses the converter, not F10)",
            },
            'workbook_date_cell_J12': norm(ws['J12'].value),
            'extracted': datetime.date.today().isoformat(),
            'method': 'openpyxl data_only=True (cached values saved by Excel), read_only=True; workbook not modified',
            'script': 'tests/ca/fixtures/extract_ref01.py',
        },
        'inputs': {k: norm(ws[c].value) for c, k in INPUT_CELLS.items()},
        'converter': {k: norm(cv[c].value) for c, k in CONVERTER_CELLS.items()},
        'summary': {k: norm(ws[c].value) for c, k in SUMMARY_CELLS.items()},
        'row_columns': {k: {'column': c, 'header_row36': headers[c], 'meaning': m} for c, k, m in ROW_COLS},
        'rows': rows,
    }
    dst = os.path.join(HERE, 'ca_ref01_workbook_saved.json')
    with open(dst, 'w') as fh:
        json.dump(out, fh, indent=1)
        fh.write('\n')
    print('wrote', dst, len(rows), 'rows; sha256', sha)


if __name__ == '__main__':
    main()
