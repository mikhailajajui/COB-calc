#!/usr/bin/env python3
"""Builds COB_Calculator_CA.xlsx -- the Canadian Cost-of-Borrowing (COB)
disclosure calculator. See docs/new-req/006-cost-of-borrowing-disclosure.md
(rewritten per docs/new-req/007-cob-canada-brd-reconciliation.md's findings)
for the spec this implements.

This REBUILD reflects the reconciled model (doc 007's ten findings), not the
prior draft:
  - payment_amount is a plain user INPUT (no PMT solve).
  - No "remaining amortization" concept -- the schedule runs from
    first_payment_date to end_date (or until closing_balance hits zero,
    whichever first). No term-vs-amortization clamp.
  - amortized_principal (P0) = loan_amount, unchanged by the fee split.
    disbursal_amount = loan_amount - financed_fees (cash fees never touch
    disbursal).
  - Interest accrues on ACTUAL CALENDAR DAYS per period (day-count
    proration), not a fixed periodic rate: period_interest = opening_balance
    x calculated_rate x (days_in_period / 365), split 365/366 across a
    leap-year boundary.
  - accrued_interest is a running carried_accrued_interest column (never
    capitalized into the opening balance).
  - Per-row payment waterfall: interest (incl. carried accrued) -> fees
    (financed + cash combined, recovered gradually) -> principal.
  - cob_amount = total_interest + ALL fees, unconditionally (no
    included_in_cob filtering).
  - APR's T = actual days from start_date to the LAST GENERATED schedule
    row's date, divided by 365 (not a nominal term_years+term_months/12).
  - Flow dropdown collapsed to 4 values (COB Requirements v2.3.docx IN-01).
  - Compounding m selected by product/rate type: fixed mortgage m=2,
    variable mortgage m=n (no conversion), personal loan (either rate
    type) m=12.

This is a separate script (not new functions appended to build_workbook.py)
but reuses that module's styling constants and helpers directly, so the two
workbooks look and behave identically at the cell-formatting level -- see
the imports below.

Run: python3 build_workbook_ca.py [output_path]
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from build_workbook import (
    FONT_NAME,
    INPUT_FONT,
    FORMULA_FONT,
    LABEL_FONT,
    NOTE_FONT,
    TITLE_FONT,
    INPUT_ROW_FILL,
    BORDER,
    MONEY_FMT,
    PCT_FMT,
    DATE_FMT,
    NUM_FMT,
    style_header_row,
    set_col_widths,
    title_block,
    _scrub_empty_string_cells,
)

SHEET_NAME_CA = "COB_CA"
SHEET_NAME_SCHEDULE = "COB_CA_Schedule"

# Percent-numbers in this workbook (e.g. 3.74 meaning 3.74%) are displayed
# with a plain high-precision NUMBER format, not Excel's "%" format -- "%"
# would multiply the stored value by 100 again, which is wrong since the
# value already IS the percent number (matches this project's
# percent-rate convention and build_workbook.py's own "0.00000000" usage
# for rate cells throughout COB_Calculator.xlsx).
RATE_FMT = "0.00000000000000"

# Bound for COB_CA_Schedule. There is no amortization-horizon concept in the
# corrected model (doc 007 finding #2) -- the schedule just runs from
# first_payment_date to end_date (or an early full payoff). Sized generously:
# 10 years at the densest supported cadence (weekly, 52/yr) = 520 rows. Same
# house convention as the rest of this workbook: extend by copying the last
# row's formulas down if a longer schedule is needed.
MAX_CA_SCHEDULE_ROWS = 520

SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12)


def _section(ws: Worksheet, row: int, text: str) -> None:
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT


def _row(ws: Worksheet, r: int, label: str, value=None, formula=None, fmt=None, is_input=False, note=None) -> str:
    """Writes one label/value (or label/formula) row and returns its value cell's
    reference ("B{r}"), the same convention as build_workbook.py's
    _labeled_inputs/_labeled_outputs but row-by-row, so every formula below can
    refer to another row by name (via the `ref` dict) instead of a hardcoded
    row number -- the exact bug class documented in docs/new-req/005 (a stale
    hardcoded column/row reference silently pointing at the wrong cell)."""
    ws.cell(row=r, column=1, value=label).font = LABEL_FONT
    cell = ws.cell(row=r, column=2, value=value if value is not None else formula)
    if is_input:
        cell.font = INPUT_FONT
        cell.fill = INPUT_ROW_FILL
        cell.border = BORDER
    else:
        cell.font = FORMULA_FONT
    if fmt:
        cell.number_format = fmt
    if note:
        nc = ws.cell(row=r, column=3, value=note)
        nc.font = NOTE_FONT
    return f"B{r}"


def xref(sheet: str, ref: str) -> str:
    """Absolute cross-sheet reference, e.g. xref('COB_CA', 'B16') -> "COB_CA!$B$16"."""
    col, row = coordinate_from_string(ref)
    return f"{sheet}!${col}${row}"


def _is_leap_expr(year_expr: str) -> str:
    """Excel has no built-in ISLEAP; standard Gregorian rule inline."""
    return f"AND(MOD({year_expr},4)=0,OR(MOD({year_expr},100)<>0,MOD({year_expr},400)=0))"


def _day_count_factor_expr(prior_expr: str, this_expr: str) -> str:
    """Equation 3's leap-year day-count split, matching the live macro's exact
    getRate/crossesLeapYear/daysInYear algorithm bit-for-bit (doc 007's
    "Addendum: VBA macro source review", verified against doc 007's full
    156-row worked vector in a standalone Python simulation before being
    ported here). Split point is JAN 1 of the following year (NOT Dec 31 of
    the prior year -- that off-by-one was this implementation's original,
    incorrect boundary and is exactly what caused a ~$0.007 total_interest
    drift against the live workbook in an earlier draft): the segment from
    prior_expr up to (but not including) Jan 1 uses prior_expr's year's own
    day-count divisor; the segment from Jan 1 onward uses this_expr's year's
    divisor. Every payment_frequency this engine supports (monthly/
    semiMonthly/biweekly/weekly) produces periods well under 365 days, so a
    period can straddle AT MOST ONE such boundary (touching exactly two
    calendar years) -- this formula handles exactly that case and is not
    generalized to a period spanning multiple full years (the macro's own
    multi-year loop case), a deliberate, documented simplification.
    """
    boundary = f"(DATE(YEAR({prior_expr})+1,1,1))"  # Jan 1 of prior_expr's year + 1
    same_year = f"({this_expr}-{prior_expr})/IF({_is_leap_expr(f'YEAR({this_expr})')},366,365)"
    cross_part1 = f"(({boundary})-{prior_expr})/IF({_is_leap_expr(f'YEAR({prior_expr})')},366,365)"
    cross_part2 = f"({this_expr}-({boundary}))/IF({_is_leap_expr(f'YEAR({this_expr})')},366,365)"
    return f"IF(YEAR({prior_expr})=YEAR({this_expr}),{same_year},{cross_part1}+{cross_part2})"


def build_cob_ca_sheet(wb: Workbook):
    ws = wb.create_sheet(SHEET_NAME_CA)
    title_block(
        ws,
        "Canadian Cost of Borrowing (COB) Disclosure -- inputs & summary",
        "Reproduces Alterna Savings' COB Calculator v7 for the same inputs (see "
        "docs/new-req/006-cost-of-borrowing-disclosure.md and "
        "docs/new-req/007-cob-canada-brd-reconciliation.md) -- a borrower-facing "
        "cost-of-borrowing ESTIMATE, not a certified regulatory disclosure engine or a "
        "general model of every Canadian lender's practice. Blue cells = edit; black = "
        "computed. Every input field stays visible for every Flow (no macros hide/show "
        "fields) -- pick Flow/Product type/Rate type below, then read the computed helper "
        "rows to see which date field, trigger-rate, and compounding rule actually apply; "
        "unused fields for your Flow are simply ignored by the formulas.",
    )

    ref = {}
    r = 4
    _section(ws, r, "Flow / product / rate-type selection")
    r += 1
    ref["flow"] = _row(
        ws, r, "Flow", value="New Mortgage/Loan", is_input=True,
        note="4 values per COB Requirements v2.3.docx IN-01 (doc 007 finding #9): New Mortgage/Loan, "
        "Renewal, Payment Change, Variable Rate Payment Change.",
    )
    r += 1
    ref["product_type"] = _row(ws, r, "Product type", value="mortgage", is_input=True)
    r += 1
    ref["rate_type"] = _row(ws, r, "Rate type", value="fixed", is_input=True)
    r += 1
    r += 1  # spacer

    ref["date_field_used"] = _row(
        ws, r, "Date field used (per spec's Flow table)",
        formula=f'=IF({ref["flow"]}="New Mortgage/Loan","Disbursal date","Renewal date")',
    )
    r += 1
    ref["carry_forward"] = _row(
        ws, r, "Accrued-interest carry-forward applies?",
        formula=f'=IF({ref["flow"]}="New Mortgage/Loan","No","Yes")',
    )
    r += 1
    ref["compounding"] = _row(
        ws, r, "Compounding convention (eq. 1 / eq. 2, docx section 4.2)",
        formula=(
            f'=IF(AND({ref["product_type"]}="mortgage",{ref["rate_type"]}="fixed"),'
            f'"Fixed mortgage: m=2, semi-annual (Interest Act s.6, CONFIRMED)",'
            f'IF(AND({ref["product_type"]}="mortgage",{ref["rate_type"]}="variable"),'
            f'"Variable mortgage: m=n, no conversion (calculated rate = contract rate)",'
            f'"Personal loan: m=12, monthly compounding"))'
        ),
    )
    r += 1
    ref["trigger_applies"] = _row(
        ws, r, "Trigger rate computed? (mortgage + variable only, all 4 flows)",
        formula=f'=IF(AND({ref["product_type"]}="mortgage",{ref["rate_type"]}="variable"),"Yes","No / N/A")',
    )
    r += 1
    r += 1  # spacer

    # Small reference lookup table (payments/year per payment_frequency), mirroring
    # PaymentFrequency!D4:E8 in the US-style workbook exactly (same 4 cadences, same
    # payments-per-year values -- reused rather than re-derived).
    ws.cell(row=4, column=4, value="Frequency").font = LABEL_FONT
    ws.cell(row=4, column=5, value="Payments/Yr").font = LABEL_FONT
    freq_table = [("monthly", 12), ("semiMonthly", 24), ("biweekly", 26), ("weekly", 52)]
    for i, (name, ppy) in enumerate(freq_table):
        ws.cell(row=5 + i, column=4, value=name).font = FORMULA_FONT
        ws.cell(row=5 + i, column=5, value=ppy).font = FORMULA_FONT
    freq_name_rng, freq_ppy_rng = "$D$5:$D$8", "$E$5:$E$8"

    _section(ws, r, "Inputs -- common to every flow")
    r += 1
    ref["loan_amount"] = _row(
        ws, r, "Loan amount ($) / current outstanding balance ($)", value=227829.65, fmt=MONEY_FMT, is_input=True,
        note="New Mortgage/Loan: the new loan/mortgage face amount, ALREADY INCLUSIVE of financed fees "
        "(IN-02). Renewal/Payment Change/Variable Rate Payment Change flows: enter the prior term's "
        "final closing_balance here instead (same field, reused per spec).",
    )
    r += 1
    ref["contract_rate"] = _row(
        ws, r, "Contract rate, nominal annual (%, e.g. 3.74 means 3.74%)", value=3.74, fmt="0.00000000",
        is_input=True,
    )
    r += 1
    ref["payment_frequency"] = _row(
        ws, r, "Payment frequency", value="weekly", is_input=True,
        note='"Accelerated Weekly" is treated identically to plain Weekly for rate/day-count purposes '
        "(BR-09) -- select weekly for either.",
    )
    r += 1
    ref["payment_amount"] = _row(
        ws, r, "Payment amount ($) -- PLAIN INPUT, never solved by this engine", value=465.46, fmt=MONEY_FMT,
        is_input=True,
        note="Doc 007 finding #1: payment_amount is user-entered in every flow, not an annuity solve. "
        "Constant across every scheduled payment.",
    )
    r += 1
    ref["first_payment_date"] = _row(ws, r, "First payment date", value=date(2026, 3, 23), fmt=DATE_FMT, is_input=True)
    r += 1
    ref["end_date"] = _row(
        ws, r, "End date (the date the schedule is run to)", value=date(2029, 3, 17), fmt=DATE_FMT, is_input=True,
        note="No amortization-horizon concept exists in this engine (doc 007 finding #2) -- the schedule "
        "commonly ends with a nonzero closing_balance still outstanding.",
    )
    r += 1
    ref["term_years"] = _row(
        ws, r, "Term years (DISPLAY ONLY -- does not drive the schedule)", value=3, fmt=NUM_FMT, is_input=True,
    )
    r += 1
    ref["term_months"] = _row(ws, r, "Term months (display only, extra months)", value=0, fmt=NUM_FMT, is_input=True)
    r += 1
    r += 1  # spacer

    _section(ws, r, "Inputs -- flow-conditional")
    r += 1
    ref["disbursal_date"] = _row(
        ws, r, "Disbursal date (New Mortgage/Loan ONLY)", value=date(2026, 3, 17), fmt=DATE_FMT, is_input=True,
    )
    r += 1
    ref["renewal_date"] = _row(
        ws, r, "Renewal date (Renewal / Payment Change / Variable Rate Payment Change ONLY)",
        value=date(2026, 3, 17), fmt=DATE_FMT, is_input=True,
    )
    r += 1
    ref["accrued_interest"] = _row(
        ws, r, "Accrued interest ($) -- Renewal/Payment Change/Variable Rate Payment Change ONLY",
        value=0, fmt=MONEY_FMT, is_input=True,
        note="NOT capitalized into the opening balance (doc 007 finding #5). Seeds the schedule's "
        "carried_accrued_interest running column instead -- see COB_CA_Schedule column F.",
    )
    r += 1
    ref["semi_annual_ref_date"] = _row(
        ws, r, "Semi-annual compounding reference date (fixed mortgages only)", value=date(2026, 3, 17), fmt=DATE_FMT,
        is_input=True,
        note="Display/reference only -- eq. 1's periodic-rate conversion does not depend on this date. "
        "Not used in any formula below.",
    )
    r += 1
    r += 1  # spacer

    _section(ws, r, "Fees")
    r += 1
    ws.cell(
        row=r, column=1,
        value=(
            "cob_amount now includes ALL fees unconditionally (financed + cash), per doc 007 finding #6. "
            "The 'Included in COB?' column below is kept for backward compatibility only -- it no longer "
            "drives any formula (do not re-wire it)."
        ),
    ).font = NOTE_FONT
    r += 2

    fee_header_row = r
    fee_headers = ["Fee name", "Amount ($)", "Financed? (Y/N)", "Included in COB? (Y/N, UNUSED)"]
    for i, h in enumerate(fee_headers, start=1):
        ws.cell(row=fee_header_row, column=i, value=h)
    style_header_row(ws, fee_header_row, 1, len(fee_headers))
    fee_first_row = fee_header_row + 1
    fee_examples = [
        ("Example financed fee (edit amount/name)", 0, "Y", "Y"),
        ("Example cash-paid fee (edit amount/name)", 0, "N", "Y"),
    ]
    fee_last_row = fee_first_row + len(fee_examples) - 1
    for i, (name, amt, fin, inc) in enumerate(fee_examples):
        rr = fee_first_row + i
        ws.cell(row=rr, column=1, value=name)
        ws.cell(row=rr, column=2, value=amt)
        ws.cell(row=rr, column=3, value=fin)
        ws.cell(row=rr, column=4, value=inc)
        for c in (1, 2, 3, 4):
            ws.cell(row=rr, column=c).font = INPUT_FONT
            ws.cell(row=rr, column=c).fill = INPUT_ROW_FILL
            ws.cell(row=rr, column=c).border = BORDER
        ws.cell(row=rr, column=2).number_format = MONEY_FMT
    fee_amt_rng = f"$B${fee_first_row}:$B${fee_last_row}"
    fee_fin_rng = f"$C${fee_first_row}:$C${fee_last_row}"
    r = fee_last_row + 2

    # --- Data validation (dropdown) lists ---
    dv_flow = DataValidation(
        type="list",
        formula1='"New Mortgage/Loan,Renewal,Payment Change,Variable Rate Payment Change"',
        allow_blank=False,
    )
    ws.add_data_validation(dv_flow)
    dv_flow.add(ws[ref["flow"]])
    dv_product = DataValidation(type="list", formula1='"mortgage,personalLoan"', allow_blank=False)
    ws.add_data_validation(dv_product)
    dv_product.add(ws[ref["product_type"]])
    dv_rate_type = DataValidation(type="list", formula1='"variable,fixed"', allow_blank=False)
    ws.add_data_validation(dv_rate_type)
    dv_rate_type.add(ws[ref["rate_type"]])
    dv_freq = DataValidation(type="list", formula1='"monthly,semiMonthly,biweekly,weekly"', allow_blank=False)
    ws.add_data_validation(dv_freq)
    dv_freq.add(ws[ref["payment_frequency"]])
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(dv_yn)
    dv_yn.add(f"C{fee_first_row}:D{fee_last_row}")

    # --- Derived / computed section ---
    _section(ws, r, "Derived quantities")
    r += 1
    ref["n"] = _row(
        ws, r, "Payments per year (n)",
        formula=f'=INDEX({freq_ppy_rng},MATCH({ref["payment_frequency"]},{freq_name_rng},0))',
        fmt=NUM_FMT,
    )
    r += 1
    ref["m"] = _row(
        ws, r, "Compounding periods/year (m, eq. 2 -- doc 007 finding #10)",
        formula=(
            f'=IF(AND({ref["product_type"]}="mortgage",{ref["rate_type"]}="fixed"),2,'
            f'IF(AND({ref["product_type"]}="mortgage",{ref["rate_type"]}="variable"),{ref["n"]},12))'
        ),
        fmt=NUM_FMT,
    )
    r += 1
    ref["calculated_rate"] = _row(
        ws, r, "Calculated rate (%, eq. 1: n*((1+rate/m)^(m/n)-1))",
        formula=f'={ref["n"]}*((1+{ref["contract_rate"]}/100/{ref["m"]})^({ref["m"]}/{ref["n"]})-1)*100',
        fmt=RATE_FMT,
        note="Unified formula for all 3 branches (m=2 fixed mortgage / m=n variable mortgage / m=12 "
        "personal loan) -- when m=n this algebraically reduces to calculated_rate = contract_rate exactly.",
    )
    r += 1
    ref["total_fees"] = _row(ws, r, "Total fees (financed + cash)", formula=f"=SUM({fee_amt_rng})", fmt=MONEY_FMT)
    r += 1
    ref["total_financed_fees"] = _row(
        ws, r, "Total financed fees", formula=f'=SUMIF({fee_fin_rng},"Y",{fee_amt_rng})', fmt=MONEY_FMT,
    )
    r += 1
    ref["total_cash_fees"] = _row(
        ws, r, "Total cash (non-financed) fees", formula=f"={ref['total_fees']}-{ref['total_financed_fees']}",
        fmt=MONEY_FMT,
    )
    r += 1
    ref["p0"] = _row(
        ws, r, "Amortized principal P0 (== loan_amount, invariant #3)",
        formula=f"={ref['loan_amount']}",
        fmt=MONEY_FMT,
        note="Doc 007 finding #4: loan_amount is ALREADY INCLUSIVE of financed fees -- financing a fee "
        "does not change this figure at all.",
    )
    r += 1
    ref["disbursal_amount"] = _row(
        ws, r, "Disbursal amount (cash actually advanced)",
        formula=f"={ref['loan_amount']}-{ref['total_financed_fees']}", fmt=MONEY_FMT,
        note="Cash (non-financed) fees do NOT reduce disbursal (IN-07, BR-04) -- only financed fees do.",
    )
    r += 1
    ref["start_date"] = _row(
        ws, r, "start_date (this flow's schedule/APR start date)",
        formula=f'=IF({ref["flow"]}="New Mortgage/Loan",{ref["disbursal_date"]},{ref["renewal_date"]})',
        fmt=DATE_FMT,
    )
    r += 1
    ref["accrued_initial"] = _row(
        ws, r, "carried_accrued_interest seed (0 for New Mortgage/Loan)",
        formula=f'=IF({ref["flow"]}="New Mortgage/Loan",0,{ref["accrued_interest"]})', fmt=MONEY_FMT,
    )
    r += 1
    ref["fees_to_recover_initial"] = _row(
        ws, r, "fees_to_recover seed (financed + cash, eq. 4)",
        formula=f"={ref['total_financed_fees']}+{ref['total_cash_fees']}", fmt=MONEY_FMT,
    )
    r += 1
    ref["trigger_rate_raw"] = _row(
        ws, r, "Trigger rate, raw/unconditional calc (eq. 6): payment*n/loan_amount*100",
        formula=f'={ref["payment_amount"]}*{ref["n"]}/{ref["p0"]}*100',
        fmt=RATE_FMT,
        note="Always computed (matches the live workbook's own 'Other related calculations!D9' cell, "
        "which has no product/rate-type gating at the formula level) -- see trigger_applies above for "
        "whether this value is actually MEANINGFUL for the current Flow/Product/Rate type. "
        "loan_amount here = P0 = the flow's current outstanding principal. Ratios stay unrounded per "
        "this project's rounding policy; the number format above truncates display only.",
    )
    r += 1
    ref["trigger_rate"] = _row(
        ws, r, "trigger_rate_percent (eq. 6 output -- mortgage + variable only, per spec)",
        formula=f'=IF({ref["trigger_applies"]}="Yes",{ref["trigger_rate_raw"]},"N/A")',
        fmt=RATE_FMT,
        note="Spec 006: null/N/A for personal loans and fixed-rate mortgages -- this is the RESTRICTED "
        "output field. See the raw/unconditional row above for the always-computed value.",
    )
    r += 1
    r += 1  # spacer

    _section(ws, r, "Outputs")
    r += 1
    return ws, ref, r, fee_first_row, fee_last_row


def _finish_cob_ca_outputs(ws: Worksheet, ref: dict, r: int, sched_ranges: dict) -> dict:
    """Second pass: writes the output rows once the Schedule sheet's (fixed) column
    layout is known, so these can reference real bounded ranges."""
    counted_rng = sched_ranges["counted"]
    date_rng = sched_ranges["date"]
    opening_rng = sched_ranges["opening"]
    interest_paid_rng = sched_ranges["interest_paid"]
    fees_paid_rng = sched_ranges["fees_paid"]
    principal_rng = sched_ranges["principal"]
    payment_rng = sched_ranges["payment"]
    closing_rng = sched_ranges["closing"]

    ref["number_of_payments"] = _row(
        ws, r, "number_of_payments", formula=f"=COUNTIF({counted_rng},TRUE)", fmt=NUM_FMT,
    )
    r += 1
    ref["final_payment_date"] = _row(
        ws, r, "final_payment_date (last row actually generated, eq. 7)",
        formula=f"=INDEX({date_rng},{ref['number_of_payments']})", fmt=DATE_FMT,
    )
    r += 1
    ref["final_closing_balance"] = _row(
        ws, r, "Final closing_balance (may be nonzero -- next flow's loan_amount input)",
        formula=f"=INDEX({closing_rng},{ref['number_of_payments']})", fmt=MONEY_FMT,
    )
    r += 1
    ref["term_days"] = _row(
        ws, r, "term_days (informational -- same day-count basis as T, doc 007 finding #7)",
        formula=f"={ref['final_payment_date']}-{ref['start_date']}", fmt=NUM_FMT,
    )
    r += 1
    ref["T"] = _row(
        ws, r, "T = term_days / 365 (eq. 7)",
        formula=f"={ref['term_days']}/365", fmt="0.00000000",
    )
    r += 1
    ref["avg_outstanding_balance"] = _row(
        ws, r, "P: average of each counted row's opening balance (eq. 7)",
        formula=(
            f"=IF({ref['number_of_payments']}=0,0,"
            f"SUMPRODUCT(({counted_rng})*({opening_rng}))/{ref['number_of_payments']})"
        ),
        fmt=MONEY_FMT,
    )
    r += 1
    ref["total_interest"] = _row(
        ws, r, "total_interest (sum of interest_paid across counted rows)",
        formula=f"=SUMIFS({interest_paid_rng},{counted_rng},TRUE)", fmt=MONEY_FMT,
    )
    r += 1
    ref["fees_recovered"] = _row(
        ws, r, "fees_recovered (sum of fees_paid, for invariant #2's reconciliation)",
        formula=f"=SUMIFS({fees_paid_rng},{counted_rng},TRUE)", fmt=MONEY_FMT,
    )
    r += 1
    ref["principal_payment"] = _row(
        ws, r, "principal_payment (sum of principal_portion across counted rows)",
        formula=f"=SUMIFS({principal_rng},{counted_rng},TRUE)", fmt=MONEY_FMT,
    )
    r += 1
    ref["total_payment"] = _row(
        ws, r, "total_payment (sum of each row's recorded payment, capped on an early full payoff row)",
        formula=f"=SUMIFS({payment_rng},{counted_rng},TRUE)", fmt=MONEY_FMT,
        note="Invariant #2: total_payment == total_interest + fees_recovered + principal_payment exactly "
        "(within floating-point rounding).",
    )
    r += 1
    ref["cob_amount"] = _row(
        ws, r, "cob_amount = total_interest + ALL fees, unconditionally (eq. 8)",
        formula=f"={ref['total_interest']}+{ref['total_financed_fees']}+{ref['total_cash_fees']}",
        fmt=MONEY_FMT,
    )
    r += 1
    ref["cob_rate"] = _row(
        ws, r, "cob_rate_percent (eq. 7 -- HARD SHORT-CIRCUIT when fees are $0, not a T/P limit)",
        formula=(
            f"=IF({ref['total_financed_fees']}+{ref['total_cash_fees']}=0,{ref['calculated_rate']},"
            f"IF(OR({ref['T']}=0,{ref['avg_outstanding_balance']}=0),\"n/a\","
            f"{ref['cob_amount']}/({ref['T']}*{ref['avg_outstanding_balance']})*100))"
        ),
        fmt=RATE_FMT,
        note="Doc 007's 'Addendum: VBA macro source review' #1 (settled via the live macro's own "
        "source): when total_financed_fees+total_cash_fees=0, cob_rate_percent = calculated_rate "
        "DIRECTLY -- this is literal macro logic (`If finFee+nonFinFee=0 Then COBRate=aRate*100`, FCPFR "
        "s.32), not an approximation, and NO choice of T/P reproduces this value as an identity from "
        "the general C/(T*P) formula -- do not remove this branch to 'simplify.' Only when fees > 0 does "
        "the general Financial Consumer Protection Framework Regulations SOR/2021-181 ss.47-48 formula "
        "apply (T plain /365 per the macro's `totalYears = termDay/365`, not leap-adjusted; P = simple "
        "unweighted average of opening balances). Unrounded per this project's rounding policy.",
    )
    r += 1
    r += 1  # spacer
    ws.cell(
        row=r, column=1,
        value=(
            "Optional: COB_Calculator_CA_macros.bas (in this folder) is STALE against this rebuild -- it "
            "hardcodes the PRIOR schedule's cell/column layout (COB_CA!B58, COB_CA!B65, an 8-column "
            "A-H schedule) which no longer matches this sheet. Treat it as legacy/unused until it is "
            "updated to match the layout below; none of the formulas on this sheet depend on it."
        ),
    ).font = NOTE_FONT
    return ref


def build_cob_ca_schedule_sheet(wb: Workbook, ref: dict):
    ws = wb.create_sheet(SHEET_NAME_SCHEDULE)
    title_block(
        ws,
        "COB_CA amortization schedule (computed) -- do not edit",
        f"One row per payment_frequency period from first_payment_date, bounded to "
        f"{MAX_CA_SCHEDULE_ROWS} rows. A row is 'counted' (column O) only if its own opening balance is "
        "still positive AND its date is on or before end_date (eq. 5, INCLUSIVE of end_date per doc 007's "
        "VBA macro source review) -- rows past that point freeze at $0/FALSE, the same self-stabilizing "
        "technique used throughout this project (see docs/new-req/005). Interest (column E) is "
        "day-count-prorated (eq. 3), matching the live macro's getRate algorithm bit-for-bit: leap-year "
        "periods are split 365/366 at the Jan-1 boundary (the segment before Jan 1 uses the PRIOR year's "
        "day-count divisor, the segment from Jan 1 onward uses the NEW year's), but this formula assumes a "
        "period never spans more than one such boundary -- true for every payment_frequency this engine "
        "supports (max ~31 days/period), but a documented simplification, not a fully general multi-year "
        "day-count routine. Values are NOT rounded to cents at each row (matching the live Alterna "
        "workbook's own full-precision behavior, confirmed "
        "in doc 007's worked vector) -- only cell display formats truncate for readability.",
    )

    headers = [
        "Period #", "Period Date", "Days in Period", "Opening Balance", "Period Interest",
        "Accrued Int. (Opening)", "Fees (Opening)", "Payment Amount", "Interest Paid", "Fees Paid",
        "Principal Portion", "Accrued Int. (Closing)", "Fees (Closing)", "Closing Balance",
        "Counted in totals?",
    ]
    header_row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, 1, len(headers))
    first_data_row = header_row + 1
    last_data_row = first_data_row + MAX_CA_SCHEDULE_ROWS - 1

    p0 = xref("COB_CA", ref["p0"])
    calculated_rate = xref("COB_CA", ref["calculated_rate"])
    payment_amount = xref("COB_CA", ref["payment_amount"])
    end_date = xref("COB_CA", ref["end_date"])
    first_payment_date = xref("COB_CA", ref["first_payment_date"])
    start_date = xref("COB_CA", ref["start_date"])
    accrued_initial = xref("COB_CA", ref["accrued_initial"])
    fees_initial = xref("COB_CA", ref["fees_to_recover_initial"])
    n_per_year = xref("COB_CA", ref["n"])

    for i in range(MAX_CA_SCHEDULE_ROWS):
        rr = first_data_row + i
        period_num = i + 1
        prev_rr = rr - 1

        ws.cell(row=rr, column=1, value=period_num)

        # Period date: explicit month/day-count stepping matching the rest of the
        # project's date arithmetic (see docs/new-req/005) -- monthly/semiMonthly use
        # whole-month rollover (semiMonthly is an approximation: 365.25/24 days/period
        # on average, a pre-existing simplification carried over from the prior draft,
        # not new to this rebuild), biweekly/weekly step by exact days.
        ws.cell(
            row=rr, column=2,
            value=(
                f'=IF({n_per_year}=12,'
                f"DATE(YEAR({first_payment_date})+INT((MONTH({first_payment_date})-1+{period_num - 1})/12),"
                f"MOD(MONTH({first_payment_date})-1+{period_num - 1},12)+1,1)+(DAY({first_payment_date})-1),"
                f'IF({n_per_year}=24,{first_payment_date}+ROUND(({period_num - 1}*365.25)/24,0),'
                f'IF({n_per_year}=26,{first_payment_date}+{(period_num - 1) * 14},'
                f'{first_payment_date}+{(period_num - 1) * 7})))'
            ),
        )
        ws.cell(row=rr, column=2).number_format = DATE_FMT

        this_date_expr = f"B{rr}"
        prior_date_expr = start_date if period_num == 1 else f"B{prev_rr}"

        ws.cell(row=rr, column=3, value=f"={this_date_expr}-{prior_date_expr}")
        ws.cell(row=rr, column=3).number_format = NUM_FMT

        opening_expr = p0 if period_num == 1 else f"N{prev_rr}"
        ws.cell(row=rr, column=4, value=f"={opening_expr}")
        ws.cell(row=rr, column=4).number_format = MONEY_FMT

        factor_expr = _day_count_factor_expr(prior_date_expr, this_date_expr)
        ws.cell(row=rr, column=5, value=f"=D{rr}*{calculated_rate}/100*({factor_expr})")
        ws.cell(row=rr, column=5).number_format = MONEY_FMT

        accrued_open_expr = accrued_initial if period_num == 1 else f"L{prev_rr}"
        ws.cell(row=rr, column=6, value=f"={accrued_open_expr}")
        ws.cell(row=rr, column=6).number_format = MONEY_FMT

        fees_open_expr = fees_initial if period_num == 1 else f"M{prev_rr}"
        ws.cell(row=rr, column=7, value=f"={fees_open_expr}")
        ws.cell(row=rr, column=7).number_format = MONEY_FMT

        # -- Waterfall (eq. 4): interest (own + carried accrued) -> fees -> principal.
        total_interest_due = f"(E{rr}+F{rr})"
        ws.cell(row=rr, column=9, value=f"=MIN({payment_amount},{total_interest_due})")  # interest_paid
        ws.cell(row=rr, column=9).number_format = MONEY_FMT

        remaining_after_interest = f"({payment_amount}-I{rr})"
        ws.cell(row=rr, column=10, value=f"=MIN({remaining_after_interest},G{rr})")  # fees_paid
        ws.cell(row=rr, column=10).number_format = MONEY_FMT

        remaining_after_fees = f"({remaining_after_interest}-J{rr})"
        ws.cell(row=rr, column=11, value=f"=MIN({remaining_after_fees},D{rr})")  # principal_portion
        ws.cell(row=rr, column=11).number_format = MONEY_FMT

        # Payment Amount (H): this flow's constant payment_amount, EXCEPT capped down
        # on an early full-payoff row (remaining_after_fees would exceed the opening
        # balance) so invariant #2's reconciliation (total_payment == total_interest +
        # fees_recovered + principal_payment) holds exactly rather than "leaking" the
        # uncollected overpayment.
        would_full_payoff = f"({remaining_after_fees}>D{rr})"
        ws.cell(row=rr, column=8, value=f"=IF({would_full_payoff},I{rr}+J{rr}+K{rr},{payment_amount})")
        ws.cell(row=rr, column=8).number_format = MONEY_FMT

        ws.cell(row=rr, column=12, value=f"=E{rr}+F{rr}-I{rr}")  # carried_accrued_interest_closing
        ws.cell(row=rr, column=12).number_format = MONEY_FMT

        ws.cell(row=rr, column=13, value=f"=G{rr}-J{rr}")  # fees_closing
        ws.cell(row=rr, column=13).number_format = MONEY_FMT

        ws.cell(row=rr, column=14, value=f"=D{rr}-K{rr}")  # closing_balance
        ws.cell(row=rr, column=14).number_format = MONEY_FMT

        # Counted (eq. 5 stop condition, INCLUSIVE of end_date -- doc 007's "Addendum:
        # VBA macro source review" #4, matching the live macro's own
        # `DateDiff("d", endDate, candidateDate) > 0` exit check): a row landing
        # exactly ON end_date IS generated/counted; only a candidate date strictly
        # AFTER end_date is excluded (hence <= here, not <). This row's own opening
        # balance must also still be positive (i.e. the prior row hadn't already
        # fully paid off). Both conditions are monotonic across the schedule (balance
        # never resurrects from 0, dates only increase), so TRUE rows always form a
        # contiguous prefix starting at row 1 -- see build_cob_ca_schedule
        # docstring/comment for why INDEX(range, number_of_payments) is valid below.
        ws.cell(row=rr, column=15, value=f"=AND(D{rr}>0,{this_date_expr}<={end_date})")

        for c in range(1, 16):
            ws.cell(row=rr, column=c).font = FORMULA_FONT

    set_col_widths(
        ws,
        {
            "A": 9, "B": 13, "C": 11, "D": 15, "E": 13, "F": 14, "G": 12, "H": 13, "I": 12, "J": 11,
            "K": 13, "L": 14, "M": 12, "N": 15, "O": 12,
        },
    )
    ws.freeze_panes = ws.cell(row=first_data_row, column=2).coordinate

    sched_ranges = {
        "date": f"COB_CA_Schedule!$B${first_data_row}:$B${last_data_row}",
        "opening": f"COB_CA_Schedule!$D${first_data_row}:$D${last_data_row}",
        "payment": f"COB_CA_Schedule!$H${first_data_row}:$H${last_data_row}",
        "interest_paid": f"COB_CA_Schedule!$I${first_data_row}:$I${last_data_row}",
        "fees_paid": f"COB_CA_Schedule!$J${first_data_row}:$J${last_data_row}",
        "principal": f"COB_CA_Schedule!$K${first_data_row}:$K${last_data_row}",
        "closing": f"COB_CA_Schedule!$N${first_data_row}:$N${last_data_row}",
        "counted": f"COB_CA_Schedule!$O${first_data_row}:$O${last_data_row}",
    }
    return ws, sched_ranges


if __name__ == "__main__":
    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "COB_Calculator_CA.xlsx"

    # Regenerating formulas into an EXISTING macro-enabled workbook (see
    # COB-xlsx/MANUAL.md, "Regenerating after macros are installed"): if the
    # target is a .xlsm that already exists, load it with keep_vba=True so its
    # vbaProject.bin (and any macro-owned sheet, e.g. COB_CA_Chart) survive
    # byte-for-byte -- openpyxl can't AUTHOR a VBA project from nothing, but it
    # can faithfully preserve one that's already there while this script tears
    # down and rebuilds just the two formula-driven sheets it owns. Any other
    # target (a fresh .xlsx, or a .xlsm that doesn't exist yet) builds a brand
    # new workbook from scratch exactly as before -- no macros, pure formulas.
    reuse_macro_workbook = out_path.suffix.lower() == ".xlsm" and out_path.exists()
    if reuse_macro_workbook:
        wb = load_workbook(out_path, keep_vba=True)
        for name in (SHEET_NAME_CA, SHEET_NAME_SCHEDULE):
            if name in wb.sheetnames:
                del wb[name]
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws_ca, ref, r_after_fees, fee_first_row, fee_last_row = build_cob_ca_sheet(wb)
    # Schedule sheet needs ref (P0/calculated_rate/payment_amount/start_date/etc.) built
    # above; COB_CA's own OUTPUT rows need the Schedule sheet's bounded ranges, which
    # need the Schedule sheet built. Two-pass: build Schedule using the partial ref dict
    # (it only needs the derived-quantity cells, already written), then come back and
    # finish COB_CA's output rows using the Schedule ranges.
    ws_sched, sched_ranges = build_cob_ca_schedule_sheet(wb, ref)
    ref = _finish_cob_ca_outputs(ws_ca, ref, r_after_fees, sched_ranges)

    set_col_widths(ws_ca, {"A": 62, "B": 18, "C": 46, "D": 14, "E": 12})

    _scrub_empty_string_cells(wb)
    wb.save(out_path)
    if reuse_macro_workbook:
        print(f"wrote {out_path} (preserved existing VBA project / macro-owned sheets)")
    else:
        print(f"wrote {out_path}")
